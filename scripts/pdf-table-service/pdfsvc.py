"""PDFの表 → Excel/CSV 変換と検算の共通部品(R1 見本・受注作業用)

抽出は PyMuPDF(fitz)、検算は別エンジンの poppler pdftotext -bbox で行う。
  - 抽出側: 各セルに「PDF上のどの矩形から取ったか(bbox)」を必ず持たせる
  - 検算側: pdftotext -bbox の単語を1文字ずつの座標に割り、
      (1) セルごと: セル矩形に入るPDFの文字の多重集合 == 出力セルの文字の多重集合(空白は無視)
      (2) 取りこぼし: 見出し領域にもどのセルにも入らないPDFの文字
      (3) 見出し: 見出し領域のPDFの文字が列名に現れるか(列名は平らにした名前なので集合で確認)
  - 結合セルを下の行へ引き継いだ値(inh=True)は、PDFに1回しか書かれていないので照合対象から外し件数だけ数える

使い方は ops/runbooks/pdf-data-service.md を参照。
"""
import csv
import hashlib
import html
import os
import re
import subprocess
from collections import Counter

import fitz  # PyMuPDF

POPPLER_PDFTOTEXT = os.environ.get(
    'POPPLER_PDFTOTEXT',
    r'C:\Users\E24054\AppData\Local\Programs\MiKTeX\miktex\bin\x64\pdftotext.exe')

WS = re.compile(r'\s+')


def strip_ws(s):
    return WS.sub('', s or '')


def cell(v, page=None, bbox=None, inh=False, verify_as=None, note=None, meta=False):
    """出力セル。verify_as: 照合時に使うPDF側の文字列(表記を変換したセルのみ)。
    meta=True: 頁番号など作業側で付けた値(PDFの文字ではないので照合しない)。"""
    return {'v': '' if v is None else str(v), 'page': page, 'bbox': tuple(bbox) if bbox else None,
            'inh': inh, 'verify_as': verify_as, 'note': note, 'meta': meta}


# ---------- poppler 側(検算用) ----------

def poppler_chars(pdf_path, workdir):
    """pdftotext -bbox の単語を1文字ずつ(ページ番号0始まり, 文字, cx, cy)に割る。"""
    out = os.path.join(workdir, os.path.basename(pdf_path) + '.bbox.html')
    subprocess.run([POPPLER_PDFTOTEXT, '-bbox', pdf_path, out], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    txt = open(out, encoding='utf8').read()
    chars = []
    for pi, pg in enumerate(re.split(r'<page ', txt)[1:]):
        for m in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>', pg):
            x0, y0, x1, y1 = map(float, m.groups()[:4])
            w = html.unescape(m.group(5))
            w = strip_ws(w)
            n = len(w)
            if n == 0:
                continue
            vertical = n > 1 and (y1 - y0) > 1.5 * (x1 - x0)
            for i, ch in enumerate(w):
                if vertical:
                    cx, cy = (x0 + x1) / 2, y0 + (i + 0.5) * (y1 - y0) / n
                else:
                    cx, cy = x0 + (i + 0.5) * (x1 - x0) / n, (y0 + y1) / 2
                chars.append((pi, ch, cx, cy))
    return chars, out


def inside(b, x, y, tol=0.6):
    return b[0] - tol <= x <= b[2] + tol and b[1] - tol <= y <= b[3] + tol


def verify(pdf_path, sheets, header_areas, workdir, notes_sheet=None):
    """sheets: [{'name','columns','rows':[[cell..]]}], header_areas: [(page, bbox)]"""
    chars, bbox_file = poppler_chars(pdf_path, workdir)
    all_sheets = list(sheets) + ([notes_sheet] if notes_sheet else [])
    targets = []  # (sheet, r, c, cell)
    inh_count = 0
    conv = []
    for sh in all_sheets:
        for r, row in enumerate(sh['rows']):
            for c, ce in enumerate(row):
                if ce.get('meta'):
                    continue
                if ce['inh']:
                    inh_count += 1
                    continue
                if ce['bbox'] is None:
                    if strip_ws(ce['v']):
                        targets.append((sh['name'], r, c, ce))  # bbox無しで値あり=照合不能として数える
                    continue
                targets.append((sh['name'], r, c, ce))
                if ce['verify_as'] is not None:
                    conv.append((sh['name'], r, c, ce))
    assigned = [None] * len(chars)
    per_cell = []
    for ti, (sn, r, c, ce) in enumerate(targets):
        got = []
        if ce['bbox'] is not None:
            for k, (pi, ch, cx, cy) in enumerate(chars):
                if pi == ce['page'] and assigned[k] is None and inside(ce['bbox'], cx, cy):
                    assigned[k] = ti
                    got.append(ch)
        want = strip_ws(ce['verify_as'] if ce['verify_as'] is not None else ce['v'])
        per_cell.append((sn, r, c, ce, want, ''.join(got)))
    mism = [(sn, r, c, ce, want, got) for (sn, r, c, ce, want, got) in per_cell
            if Counter(want) != Counter(got) or ce['bbox'] is None]
    # 見出し領域
    header_chars = [(pi, ch) for k, (pi, ch, cx, cy) in enumerate(chars)
                    if assigned[k] is None and any(pi == hp and inside(hb, cx, cy) for hp, hb in header_areas)]
    header_idx = set(k for k, (pi, ch, cx, cy) in enumerate(chars)
                     if assigned[k] is None and any(pi == hp and inside(hb, cx, cy) for hp, hb in header_areas))
    unassigned = [(pi, ch, round(cx), round(cy)) for k, (pi, ch, cx, cy) in enumerate(chars)
                  if assigned[k] is None and k not in header_idx]
    colchars = set(''.join(strip_ws(col) for sh in sheets for col in sh['columns']))
    header_missing = sorted(set(ch for pi, ch in header_chars if ch not in colchars))
    # ページ別集計
    pages = sorted(set(pi for pi, *_ in chars))
    per_page = []
    for p in pages:
        n_pdf = sum(1 for (pi, *_rest) in chars if pi == p)
        n_hdr = sum(1 for (pi, ch) in header_chars if pi == p)
        cells_p = [x for x in per_cell if x[3]['page'] == p]
        n_cells = len(cells_p)
        n_bad = sum(1 for x in mism if x[3]['page'] == p)
        n_un = sum(1 for x in unassigned if x[0] == p)
        n_matched_chars = sum(len(x[5]) for x in cells_p)
        per_page.append({'page': p + 1, 'pdf_chars': n_pdf, 'header_chars': n_hdr, 'cell_chars': n_matched_chars,
                         'cells': n_cells, 'cell_mismatch': n_bad, 'unassigned': n_un})
    return {'bbox_file': bbox_file, 'per_page': per_page, 'mismatch': mism, 'unassigned': unassigned,
            'header_missing': header_missing, 'inherited': inh_count, 'converted': conv,
            'total_cells': len(per_cell)}


# ---------- 書き出し ----------

def render_pages(pdf_path, outdir, dpi=100, prefix='page'):
    d = fitz.open(pdf_path)
    names = []
    for i, p in enumerate(d):
        n = f'{prefix}-{i + 1}.png'
        p.get_pixmap(dpi=dpi).save(os.path.join(outdir, n))
        names.append(n)
    return names


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        h.update(f.read())
    return h.hexdigest()


def write_xlsx(path, sheets, notes_sheet, source_lines):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    hdr_fill = PatternFill('solid', fgColor='DDEBF7')
    inh_font = Font(color='808080', italic=True)
    for sh in sheets + [notes_sheet]:
        ws = wb.create_sheet(sh['name'][:31])
        ws.append(sh['columns'])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical='top')
        for row in sh['rows']:
            vals = []
            for ce in row:
                v = ce['v']
                vals.append(int(v) if re.fullmatch(r'\d+', v or '') and sh.get('numeric_ok', True) else v)
            ws.append(vals)
            for j, ce in enumerate(row):
                x = ws.cell(row=ws.max_row, column=j + 1)
                x.alignment = Alignment(wrap_text=True, vertical='top')
                if ce['inh']:
                    x.font = inh_font
        ws.freeze_panes = 'A2'
        for j, col in enumerate(sh['columns']):
            longest = max([len(strip_ws(col))] + [max((len(t) for t in r[j]['v'].split('\n')), default=0) for r in sh['rows'] if j < len(r)])
            ws.column_dimensions[get_column_letter(j + 1)].width = min(max(6, longest * 2 + 2), 60)
    ws = wb.create_sheet('出典と凡例')
    for line in source_lines:
        ws.append([line])
    ws.column_dimensions['A'].width = 120
    wb.save(path)


def write_csvs(outdir, sheets):
    names = []
    for sh in sheets:
        n = f"{sh['name']}.csv"
        with open(os.path.join(outdir, n), 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            w.writerow(sh['columns'])
            for row in sh['rows']:
                w.writerow([ce['v'] for ce in row])
        names.append(n)
    return names


def outside_blocks(page, table_bboxes, page_index):
    """表の外にある文字ブロック(題名・注記)を 注記シートの行にする。"""
    rows = []
    for b in page.get_text('blocks'):
        x0, y0, x1, y1, t = b[0], b[1], b[2], b[3], b[4]
        if not strip_ws(t):
            continue
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        if any(inside(tb, cx, cy, tol=0) for tb in table_bboxes):
            continue
        rows.append([cell(str(page_index + 1), meta=True), cell(t.strip(), page_index, (x0, y0, x1, y1))])
    return rows


def report_md(res, max_list=30):
    """検算結果の表(Markdown)を返す。"""
    L = []
    L.append('| 頁 | PDFの文字数(pdftotext) | うち見出し領域 | セルに入った文字 | 照合したセル数 | 食い違いセル | どのセルにも入らない文字 |')
    L.append('|---:|---:|---:|---:|---:|---:|---:|')
    for p in res['per_page']:
        L.append(f"| {p['page']} | {p['pdf_chars']} | {p['header_chars']} | {p['cell_chars']} | {p['cells']} | {p['cell_mismatch']} | {p['unassigned']} |")
    tot = lambda k: sum(p[k] for p in res['per_page'])
    L.append(f"| 計 | {tot('pdf_chars')} | {tot('header_chars')} | {tot('cell_chars')} | {tot('cells')} | {tot('cell_mismatch')} | {tot('unassigned')} |")
    return '\n'.join(L)
